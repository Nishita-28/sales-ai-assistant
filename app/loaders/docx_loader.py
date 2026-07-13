"""DOCX document extraction and parsing logic."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from app.loaders.shared import _build_table_block, _clean_text


def _cell_is_bold_docx(cell) -> bool:
    runs = [r for p in cell.paragraphs for r in p.runs if r.text.strip()]
    if not runs:
        return False
    return all(r.bold for r in runs)


def resolve_table_grid(table) -> list[list[str]]:
    """docx: merges are already text-duplicated by python-docx's row.cells."""
    grid = [[_clean_text(cell.text) for cell in row.cells] for row in table.rows]
    max_len = max((len(row) for row in grid), default=0)
    for row in grid:
        row.extend([""] * (max_len - len(row)))
    return grid


def _paragraph_type_docx(paragraph) -> str:
    style_name = ""
    if paragraph.style is not None and paragraph.style.name:
        style_name = paragraph.style.name.lower()
    if style_name.startswith("heading") or style_name == "title":
        return "heading"
    if "list" in style_name:
        return "list_item"

    text = paragraph.text.strip()
    runs = [r for r in paragraph.runs if r.text.strip()]
    if text and runs and len(text) <= 100:
        all_bold = all(r.bold for r in runs)
        sizes = [r.font.size.pt 
        for r in runs if r.font.size is not None]
        looks_larger = bool(sizes) and min(sizes) >= 13
        if all_bold or looks_larger:
            return "heading"
    return "paragraph"


def extract_docx_blocks(file_path: str | Path) -> list[dict[str, Any]]:
    from docx import Document
    from docx.text.paragraph import Paragraph

    path = Path(file_path)
    filename = path.name
    document = Document(path)

    blocks: list[dict[str, Any]] = []
    table_index = 0
    current_section = ""
    block_index = 0

    for element in document.element.body.iterchildren():
        tag = element.tag.split("}")[-1]

        if tag == "p":
            paragraph = Paragraph(element, document)
            text = _clean_text(paragraph.text)
            if not text:
                continue
            block_type = _paragraph_type_docx(paragraph)
            if block_type == "heading":
                current_section = text
            blocks.append({
                "type": block_type,
                "text": text,
                "metadata": {"filename": filename, "block_index": block_index, "section_title": current_section or None},
            })
            block_index += 1

        elif tag == "tbl":
            table = document.tables[table_index]
            grid = resolve_table_grid(table)
            is_bold_fn = lambda r, c, t=table: _cell_is_bold_docx(t.rows[r].cells[c])
            block = _build_table_block(grid, is_bold_fn, table_index, current_section)
            block["metadata"] = {
                "filename": filename, "block_index": block_index, "section_title": current_section or None,
                "table_index": table_index, "num_rows": len(grid), "num_columns": len(grid[0]) if grid else 0,
            }
            blocks.append(block)
            table_index += 1
            block_index += 1

    return blocks
