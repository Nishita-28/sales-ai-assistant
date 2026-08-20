"""Sales Aid page: generates a short, customer-facing comparison document
for a specific use case, grounded in the approved knowledge base -- meant
to be handed to a customer's technical champion to circulate internally.
"""
from __future__ import annotations

import io
import re

import streamlit as st
from docx import Document

from app import theme
from app.deal_picker import render_deal_picker, set_active_deal
from app.deals_store import update_deal_fields
from app.product_index import load_product_index
from app.response_generator import ResponseGeneratorError
from app.retriever import RetrieverError
from app.sales_aid_generator import finalize_sales_aid, stream_sales_aid
from app.sales_aid_store import list_sales_aids_for_deal, record_sales_aid

EXAMPLE_USE_CASE = (
    "Multiple potential leak spots in close vicinity, with 100% H2 contained in pipelines."
)


def render_sales_aid_page() -> None:
    st.title("Sales Aid Generator")
    st.caption(
        "Describe the customer's use case to get a short comparison document you can hand to "
        "the customer's technical champion -- grounded in the approved knowledge base only."
    )

    deal_id, customer_name, company, deal_use_case = render_deal_picker("sales_aid")

    # Reset on deal switch so a previous deal's generated sales aid doesn't
    # linger on screen mislabeled as belonging to the new one.
    if st.session_state.get("active_deal_id") != deal_id:
        set_active_deal(deal_id)
        st.session_state.pop("sales_aid_result", None)

    # Everything below needs an active deal to attach to.
    if deal_id is None:
        st.info("Pick an existing deal, or start a new one above (Customer Name + Company), to continue.")
        return

    try:
        products, _ = load_product_index()
    except (FileNotFoundError, OSError):
        products = []
    product_names = sorted({p.product_name for p in products})

    # st.form so Enter submits the same as clicking the button. Keyed
    # per-deal so switching deals loads that deal's own saved use case.
    with st.form(f"sales_aid_form_{deal_id}"):
        use_case = st.text_area("Use case / scenario", value=deal_use_case, placeholder=EXAMPLE_USE_CASE, height=100)
        mnst_products = st.multiselect(
            "MNST product(s) to feature (optional)",
            options=product_names,
            help="Leave blank to let the AI pick whichever product(s) in the approved documents "
            "best fit the use case -- pick one or more here to lock the comparison to exactly "
            "those products instead.",
        )
        compare_against = st.text_input(
            "Compare against (optional)",
            placeholder="Leave blank to let the AI pick",
        )
        generate = st.form_submit_button("Generate sales aid", type="primary")

    if generate:
        if not use_case.strip():
            st.error("Describe the use case first.")
        else:
            try:
                with st.spinner("Checking approved documents..."):
                    matches, text_stream = stream_sales_aid(use_case, compare_against, mnst_products=mnst_products)
                with st.container(border=theme.is_enterprise_theme()):
                    raw_reply = st.write_stream(text_stream)
            except (RetrieverError, ResponseGeneratorError) as e:
                st.error(f"Something went wrong generating the sales aid: {e}")
            else:
                # deal_id is guaranteed set here; the function returns early above when it's None.
                update_deal_fields(deal_id, use_case=use_case)
                set_active_deal(deal_id)

                result = finalize_sales_aid(use_case, matches, raw_reply)
                st.session_state.sales_aid_result = result
                record_sales_aid(use_case, compare_against, result.to_dict(), deal_id=deal_id)
                st.rerun()

    past = list_sales_aids_for_deal(deal_id)
    if past:
        with st.expander(f"{len(past)} previously generated for this deal"):
            for row in reversed(past):
                st.caption(f"{row['created_at']} -- {row['use_case'][:60]}")

    result = st.session_state.get("sales_aid_result")
    if not result:
        return

    st.divider()

    if result.ready_for_customer:
        st.success("Ready to share with the customer.")
    else:
        st.warning(
            "Internal draft -- needs review before sharing with the customer "
            f"(risk category: {result.risk})."
        )

    if result.customer_priorities:
        st.markdown("**Customer priorities identified**")
        for p in result.customer_priorities:
            st.markdown(f"- {p}")

    with st.container(border=theme.is_enterprise_theme()):
        if result.title:
            st.subheader(result.title)
        if result.use_case_framing:
            st.write(result.use_case_framing)

        if result.comparison:
            st.markdown("**Comparison**")
            st.markdown("\n".join(result.comparison))

    st.text_area("Customer-ready summary", value=result.customer_summary, height=150)

    if result.sources:
        with st.expander("Drawn from"):
            for name, _ in result.sources:
                st.markdown(f"- `{name}`")

    st.download_button(
        "Download sales aid (.docx)",
        data=_sales_aid_as_docx(result),
        file_name=f"{(result.title or 'sales_aid').strip().replace(' ', '_')}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


_TABLE_SEPARATOR_RE = re.compile(r":?-+:?")


def _parse_markdown_table(lines: list[str]) -> list[list[str]]:
    """Converts the LLM's raw pipe-delimited markdown table (result.comparison)
    into plain rows of cell text, dropping the "|---|---|" separator row, so it
    can be rebuilt as a real docx table."""
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(_TABLE_SEPARATOR_RE.fullmatch(c) for c in cells):
            continue
        rows.append(cells)
    return rows


def _sales_aid_as_docx(result) -> bytes:
    """Word export of the sales aid, with the comparison as a real docx table
    for correct alignment in Word/Google Docs/Outlook."""
    doc = Document()
    doc.add_heading(result.title or "Sales Aid", level=1)

    if result.customer_priorities:
        doc.add_heading("Customer priorities identified", level=2)
        for p in result.customer_priorities:
            doc.add_paragraph(p, style="List Bullet")

    if result.use_case_framing:
        doc.add_paragraph(result.use_case_framing)

    rows = _parse_markdown_table(result.comparison)
    if rows:
        doc.add_heading("Comparison", level=2)
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Light Grid Accent 1"
        for r, row_cells in enumerate(rows):
            for c, text in enumerate(row_cells[:cols]):
                cell = table.cell(r, c)
                cell.text = text
                if r == 0:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True

    if result.customer_summary:
        doc.add_heading("Customer-ready summary", level=2)
        doc.add_paragraph(result.customer_summary)

    if result.sources:
        doc.add_heading("Drawn from", level=2)
        for name, _ in result.sources:
            doc.add_paragraph(name, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
